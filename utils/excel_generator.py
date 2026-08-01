from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def generate_attendance_excel(title, records, metadata=None):
    """
    Generates an Excel spreadsheet (.xlsx) for attendance records.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Sheet"

    # Styling definitions
    title_font = Font(name="Calibri", size=16, bold=True, color="1E293B")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    meta_font = Font(name="Calibri", size=10, bold=True, color="475569")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    row = 1
    ws.cell(row=row, column=1, value=title).font = title_font
    row += 2

    if metadata:
        for k, v in metadata.items():
            ws.cell(row=row, column=1, value=f"{k}:").font = meta_font
            ws.cell(row=row, column=2, value=str(v)).alignment = left_align
            row += 1
        row += 1

    headers = ["S.No", "Roll Number", "Enrollment No", "Student Name", "Subject Code", "Subject Name", "Date", "Marked Time", "Status", "Method"]
    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_num, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    
    ws.row_dimensions[row].height = 25
    header_row = row

    for idx, r in enumerate(records, 1):
        row += 1
        data = [
            idx,
            r.student.roll_number,
            r.student.enrollment_number,
            r.student.user.get_full_name_or_username(),
            r.subject.code,
            r.subject.name,
            r.date.strftime("%Y-%m-%d"),
            r.marked_at.strftime("%H:%M:%S"),
            r.get_status_display(),
            r.get_verification_method_display()
        ]
        
        for col_num, val in enumerate(data, 1):
            cell = ws.cell(row=row, column=col_num, value=val)
            cell.alignment = center_align if col_num in [1, 7, 8, 9, 10] else left_align
            cell.border = thin_border

    # Adjust column widths automatically
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
